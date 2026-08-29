#!/usr/bin/env python
# -*- coding: utf-8 -*-
import openpyxl
import sys

try:
    wb = openpyxl.load_workbook('Unjumble Game 피드백.xlsx')
    print("=" * 100)
    print(f"Excel Sheets: {wb.sheetnames}")
    print("=" * 100)
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f"\n\n>>> SHEET: {sheet_name} <<<\n")
        
        row_count = 0
        for row in ws.iter_rows(values_only=True):
            row_count += 1
            if row_count > 60:
                print(f"\n... (More rows exist)")
                break
            
            if any(cell is not None for cell in row):
                text = " | ".join(str(cell)[:80] if cell is not None else "" for cell in row)
                print(text)
    
    print("\n" + "=" * 100)
    print("Feedback file loaded successfully!")
    
except Exception as e:
    print(f"Error: {e}")
    import traceback
    traceback.print_exc()
